"""
Export Service — CSV, JSON, PDF, Excel report generation
"""

import os
import csv
import json
from datetime import datetime
from typing import List
from models import OptimizationRun, OptimizedRoute, Delivery
from services.optimization_service import get_run_routes
from rich.console import Console

console = Console()
EXPORTS_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "exports")
os.makedirs(EXPORTS_DIR, exist_ok=True)


def _ts() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def export_routes_csv(run: OptimizationRun) -> str:
    routes = get_run_routes(run.id)
    filepath = os.path.join(EXPORTS_DIR, f"route_run{run.id}_{_ts()}.csv")
    with open(filepath, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=[
            "stop_order", "delivery_id", "customer_name", "address",
            "latitude", "longitude", "distance_from_prev_km"
        ])
        writer.writeheader()
        for r in routes:
            writer.writerow({
                "stop_order": r.stop_order + 1,
                "delivery_id": r.delivery_id,
                "customer_name": r.customer_name,
                "address": r.address,
                "latitude": r.latitude,
                "longitude": r.longitude,
                "distance_from_prev_km": round(r.distance_from_prev_km or 0, 4),
            })
    console.print(f"[green]✓ CSV exported: {filepath}[/green]")
    return filepath


def export_routes_json(run: OptimizationRun) -> str:
    routes = get_run_routes(run.id)
    filepath = os.path.join(EXPORTS_DIR, f"route_run{run.id}_{_ts()}.json")
    data = {
        "run_id": run.id,
        "run_name": run.run_name,
        "total_distance_km": round(run.best_distance_km or 0, 4),
        "execution_time_sec": round(run.execution_time_sec or 0, 2),
        "generations": run.generations_completed,
        "created_at": run.created_at.isoformat() if run.created_at else "",
        "route": [
            {
                "stop": r.stop_order + 1,
                "delivery_id": r.delivery_id,
                "customer_name": r.customer_name,
                "address": r.address,
                "latitude": r.latitude,
                "longitude": r.longitude,
                "distance_from_prev_km": round(r.distance_from_prev_km or 0, 4),
            }
            for r in routes
        ],
    }
    with open(filepath, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)
    console.print(f"[green]✓ JSON exported: {filepath}[/green]")
    return filepath


def export_deliveries_csv(deliveries: List[Delivery]) -> str:
    filepath = os.path.join(EXPORTS_DIR, f"deliveries_{_ts()}.csv")
    with open(filepath, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=[
            "delivery_id", "customer_name", "address",
            "latitude", "longitude", "contact_number", "notes", "created_at"
        ])
        writer.writeheader()
        for d in deliveries:
            writer.writerow({
                "delivery_id": d.delivery_id,
                "customer_name": d.customer_name,
                "address": d.address,
                "latitude": d.latitude or "",
                "longitude": d.longitude or "",
                "contact_number": d.contact_number or "",
                "notes": d.notes or "",
                "created_at": d.created_at.isoformat() if d.created_at else "",
            })
    console.print(f"[green]✓ Deliveries exported: {filepath}[/green]")
    return filepath


def export_history_csv(runs: List[OptimizationRun]) -> str:
    filepath = os.path.join(EXPORTS_DIR, f"history_{_ts()}.csv")
    with open(filepath, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=[
            "run_id", "run_name", "num_deliveries", "best_distance_km",
            "execution_time_sec", "generations_completed", "population_size",
            "mutation_rate", "crossover_rate", "created_at"
        ])
        writer.writeheader()
        for r in runs:
            writer.writerow({
                "run_id": r.id,
                "run_name": r.run_name,
                "num_deliveries": r.num_deliveries,
                "best_distance_km": round(r.best_distance_km or 0, 4),
                "execution_time_sec": round(r.execution_time_sec or 0, 2),
                "generations_completed": r.generations_completed,
                "population_size": r.population_size,
                "mutation_rate": r.mutation_rate,
                "crossover_rate": r.crossover_rate,
                "created_at": r.created_at.isoformat() if r.created_at else "",
            })
    return filepath


def export_routes_pdf(run: OptimizationRun) -> str:
    """Generate a formatted PDF report using ReportLab."""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.units import cm
    except ImportError:
        console.print("[red]reportlab not installed. Run: pip install reportlab[/red]")
        return ""

    routes = get_run_routes(run.id)
    filepath = os.path.join(EXPORTS_DIR, f"report_run{run.id}_{_ts()}.pdf")
    doc = SimpleDocTemplate(filepath, pagesize=A4, topMargin=2 * cm, bottomMargin=2 * cm)
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle("title", fontSize=18, spaceAfter=12, textColor=colors.HexColor("#1a237e"))
    header_style = ParagraphStyle("header", fontSize=12, spaceAfter=6, textColor=colors.HexColor("#333"))
    body_style = styles["BodyText"]

    elements = []
    elements.append(Paragraph("🚚 Last-Mile Delivery Route Optimizer", title_style))
    elements.append(Paragraph(f"Optimization Report — Run #{run.id}", header_style))
    elements.append(Spacer(1, 0.4 * cm))

    # Summary table
    summary_data = [
        ["Parameter", "Value"],
        ["Run Name", run.run_name or "—"],
        ["Total Distance", f"{run.best_distance_km:.3f} km"],
        ["Stops", str(run.num_deliveries)],
        ["Execution Time", f"{run.execution_time_sec:.2f} sec"],
        ["Generations", str(run.generations_completed)],
        ["Population Size", str(run.population_size)],
        ["Mutation Rate", str(run.mutation_rate)],
        ["Date", run.created_at.strftime("%Y-%m-%d %H:%M") if run.created_at else "—"],
    ]
    summary_table = Table(summary_data, colWidths=[7 * cm, 10 * cm])
    summary_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1a237e")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("BACKGROUND", (0, 1), (-1, -1), colors.HexColor("#f5f5f5")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#e8eaf6")]),
    ]))
    elements.append(summary_table)
    elements.append(Spacer(1, 0.6 * cm))
    elements.append(Paragraph("Optimized Route Sequence", header_style))

    route_data = [["#", "Delivery ID", "Customer", "Address", "Dist (km)"]]
    for r in routes:
        route_data.append([
            str(r.stop_order + 1),
            r.delivery_id,
            r.customer_name[:20],
            r.address[:40],
            f"{r.distance_from_prev_km:.3f}" if r.distance_from_prev_km else "—",
        ])
    route_table = Table(route_data, colWidths=[1 * cm, 2.5 * cm, 4 * cm, 7 * cm, 2 * cm])
    route_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#283593")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.lightgrey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#e8eaf6")]),
    ]))
    elements.append(route_table)

    doc.build(elements)
    console.print(f"[green]✓ PDF exported: {filepath}[/green]")
    return filepath


def export_excel(run: OptimizationRun) -> str:
    """Export route to Excel using openpyxl."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        console.print("[red]openpyxl not installed.[/red]")
        return ""

    routes = get_run_routes(run.id)
    filepath = os.path.join(EXPORTS_DIR, f"route_run{run.id}_{_ts()}.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Optimized Route"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1A237E")
    headers = ["Stop #", "Delivery ID", "Customer Name", "Address", "Latitude", "Longitude", "Dist from Prev (km)"]

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for r in routes:
        ws.append([
            r.stop_order + 1, r.delivery_id, r.customer_name,
            r.address, r.latitude, r.longitude,
            round(r.distance_from_prev_km or 0, 4)
        ])

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    wb.save(filepath)
    console.print(f"[green]✓ Excel exported: {filepath}[/green]")
    return filepath
